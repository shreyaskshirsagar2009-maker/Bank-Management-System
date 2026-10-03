import openpyxl
import os

file = "bank.xlsx"


# Create Excel file
if not os.path.exists(file):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Account Number", "Name", "Mobile", "Age", "Balance"])
    wb.save(file)


def menu():
    print("\nWelcome to the bank system")
    print("1. Create account.")
    print("2. Deposite money")
    print("3. Widrow")
    print("4. View bal")
    print("5. Exit")

    ch = int(input("Enter your choice: "))

    if ch == 1:
        acc()

    elif ch == 2:
        deposite()

    elif ch == 3:
        widrow()

    elif ch == 4:
        viewbal()

    elif ch == 5:
        print("Thank you")

    else:
        print("Enter valid choice")
        menu()


def acc():
    print("\n----- Create Account -----")

    account = input("Enter your account number: ")
    name = input("Enter your name: ")
    mobile = input("Enter your mobile number: ")
    age = int(input("Enter your age: "))
    balance = float(input("Enter your balance: "))

    wb = openpyxl.load_workbook(file)
    ws = wb.active

    ws.append([account, name, mobile, age, balance])

    wb.save(file)

    print("Account created successfully!")

    menu()


def deposite():
    print("\n----- Deposite Money -----")

    account = input("Enter your account number: ")
    amount = float(input("Enter money to deposite: "))

    wb = openpyxl.load_workbook(file)
    ws = wb.active

    for row in range(2, ws.max_row + 1):

        if str(ws.cell(row, 1).value) == account:

            balance = float(ws.cell(row, 5).value)

            balance = balance + amount

            ws.cell(row, 5).value = balance

            wb.save(file)

            print("Money deposited successfully")
            print("Current balance:", balance)

            menu()
            return

    print("Account not found")
    menu()


def widrow():
    print("\n----- Widrow Money -----")

    account = input("Enter your account number: ")
    amount = float(input("Enter money to widrow: "))

    wb = openpyxl.load_workbook(file)
    ws = wb.active

    for row in range(2, ws.max_row + 1):

        if str(ws.cell(row, 1).value) == account:

            balance = float(ws.cell(row, 5).value)

            if amount > balance:
                print("You do not have sufficient money")
                print("Your current balance is:", balance)

            else:
                balance = balance - amount

                ws.cell(row, 5).value = balance

                wb.save(file)

                print("Money withdrawn successfully")
                print("Remaining balance:", balance)

            menu()
            return

    print("Account not found")
    menu()


def viewbal():
    print("\n----- View Balance -----")

    account = input("Enter your account number: ")

    wb = openpyxl.load_workbook(file)
    ws = wb.active

    for row in range(2, ws.max_row + 1):

        if str(ws.cell(row, 1).value) == account:

            name = ws.cell(row, 2).value
            balance = ws.cell(row, 5).value

            print("Name:", name)
            print("Current balance:", balance)

            menu()
            return

    print("Account not found")
    menu()


menu()
